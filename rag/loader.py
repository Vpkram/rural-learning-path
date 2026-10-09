"""Extract local PDF and DOCX documents while retaining available page data."""

from __future__ import annotations

from collections.abc import Iterator
from pathlib import Path

from rag.models import DocumentPage


class DocumentLoadError(ValueError):
    """Raised when a supported document cannot be read or has no text."""


def document_page_count(path: str | Path) -> int:
    """Return the page count for a PDF or one logical page for a DOCX."""
    document_path = Path(path)
    if not document_path.is_file():
        raise DocumentLoadError(f"Document does not exist or is not a file: {document_path}")
    suffix = document_path.suffix.casefold()
    try:
        if suffix == ".pdf":
            import fitz

            with fitz.open(document_path) as document:
                page_count = getattr(document, "page_count", None)
                if isinstance(page_count, int):
                    return page_count
                try:
                    return len(document)
                except TypeError:
                    return sum(1 for _page in document)
        if suffix == ".docx":
            from docx import Document

            Document(document_path)
            return 1
    except Exception as exc:
        raise DocumentLoadError(f"Could not inspect {document_path.name}: {exc}") from exc
    raise DocumentLoadError("Only PDF and DOCX textbook files are supported.")


def iter_document_pages(
    path: str | Path,
    *,
    start_page: int = 1,
    end_page: int | None = None,
) -> Iterator[DocumentPage]:
    """Yield extracted pages incrementally, keeping PDF page numbers one-based.

    PDF page numbers are one-based. DOCX files do not reliably store page
    boundaries, so their extracted text uses page_number=None.
    """
    document_path = Path(path)
    if not document_path.is_file():
        raise DocumentLoadError(f"Document does not exist or is not a file: {document_path}")

    suffix = document_path.suffix.casefold()
    if suffix not in {".pdf", ".docx"}:
        raise DocumentLoadError("Only PDF and DOCX textbook files are supported.")
    if isinstance(start_page, bool) or not isinstance(start_page, int) or start_page < 1:
        raise ValueError("start_page must be a positive integer.")
    if end_page is not None and (
        isinstance(end_page, bool) or not isinstance(end_page, int) or end_page < start_page
    ):
        raise ValueError("end_page must be an integer greater than or equal to start_page.")

    try:
        if suffix == ".pdf":
            import fitz

            with fitz.open(document_path) as document:
                page_count = getattr(document, "page_count", None)
                if not isinstance(page_count, int):
                    try:
                        page_count = len(document)
                    except TypeError:
                        page_count = None
                if page_count is not None and start_page > page_count:
                    raise ValueError("start_page is past the end of the PDF.")
                for page_index, page in enumerate(document):
                    page_number = page_index + 1
                    if page_number < start_page:
                        continue
                    if end_page is not None and page_number > end_page:
                        break
                    text = page.get_text("text")
                    if text and text.strip():
                        yield DocumentPage(
                            source=document_path.name,
                            page_number=page_number,
                            text=text,
                        )
        else:
            if start_page != 1 or end_page not in (None, 1):
                raise ValueError("DOCX documents support only page 1 because they have no fixed page boundaries.")
            from docx import Document

            document = Document(document_path)
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            text = "\n\n".join(parts)
            if text.strip():
                yield DocumentPage(source=document_path.name, page_number=None, text=text)
    except DocumentLoadError:
        raise
    except ValueError:
        raise
    except Exception as exc:
        raise DocumentLoadError(f"Could not extract text from {document_path.name}: {exc}") from exc


def load_document(
    path: str | Path,
    *,
    start_page: int = 1,
    end_page: int | None = None,
) -> list[DocumentPage]:
    """Extract selected non-empty pages as a list for existing callers."""
    pages = list(iter_document_pages(path, start_page=start_page, end_page=end_page))
    if not pages:
        raise DocumentLoadError(f"No readable text was found in {Path(path).name}.")
    return pages
